from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
w=Path(__file__).parent
p=w/'audit_replay_threshold_station_20260918/Documents/CG_Supplementary_ReplayVerified.docx'
d=Document(p)
active=False
for para in d.paragraphs:
    if 'Historical roles' in para.text or 'Historical data roles' in para.text:
        print('heading',repr(para.text));active=True
    if para.text.startswith('Supplementary Table S23.'):
        active=True
    if active:
        for run in para.runs:
            run.font.name='Times New Roman';run.font.size=Pt(10)
            rf=run._element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                rf.attrib.pop(qn('w:'+attr),None)
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            if 'Fresh v2 replay: all 106' in cell.text:
                for para in cell.paragraphs:
                    for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(9)
d.save(p)
