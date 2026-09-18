from pathlib import Path
import csv,json,re,hashlib
from docx import Document
from lxml import etree
w=Path(__file__).parent;r=w/'audit_replay_threshold_station_20260918';docs=r/'Documents'
source=Path(r'C:\Drive E\Publikasi\2026-zd Computer & Geosciences\Submission\CG_Revision_20260918_SourceVerified\Documents')
results={}
for kind in ['MainText','Supplementary','Cover_Letter','Highlights']:
 d=Document(docs/f'CG_{kind}_ReplayVerified.docx');old=Document(source/f'CG_{kind}_SourceVerified.docx')
 text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
 for bad in ['No current model-inference replay is claimed','A fresh model-inference replay is not claimed','does not rerun current model inference','uniqueness was not exhaustively assessed','221 unseen sources']:
  assert bad not in text,(kind,bad)
 assert [etree.tostring(e, method="c14n") for e in d._element.xpath('.//m:oMath')]==[etree.tostring(e, method="c14n") for e in old._element.xpath('.//m:oMath')]
 if kind=='MainText':
  assert len(d.inline_shapes)==14 and len(d.tables)==7
  abstract=next(p.text for p in d.paragraphs if p.text.startswith('Source overlap can compromise'))
  assert len(abstract.split())<=250
  assert 'Supplementary Table S9' in text
 if kind=='Supplementary':
  assert len(d.inline_shapes)==7 and len(d.tables)==27
  for i in range(1,28):assert re.search(r'Table S'+str(i)+r'\.',text)
  b=list(csv.DictReader((r/'threshold/threshold_paired_bootstrap.csv').open()))
  labels={'Stack Full3C − Z':'Stack Full3C minus Zonly','PhaseNet Full3C − Z':'PhaseNet Full3C minus Zonly','Stack − PhaseNet Full3C':'Stack minus PhaseNet Full3C','Stack − PhaseNet Z':'Stack minus PhaseNet Zonly'}
  for row in d.tables[26].rows[1:]:
   a=[c.text for c in row.cells];policy,phase=a[0].split(' / ');rr=[x for x in b if x['Policy']==policy and x['Phase']==phase and x['Comparison']==labels[a[1]]]
   assert a[2]==f"{min(float(x['Delta']) for x in rr):.4f}" and a[3]==f"{max(float(x['Delta']) for x in rr):.4f}"
   assert a[4]==f"{sum(float(x['CI_low'])>0 for x in rr)} / {sum(float(x['CI_high'])<0 for x in rr)} of 27"
 if kind=='Highlights':
  ps=[p for p in d.paragraphs if p.text and p.text!='Highlights'];assert len(ps)==5 and all(len(p.text)<=85 for p in ps)
  for p in ps:assert p._p.xpath('./w:pPr/w:numPr') or p.style.name.startswith('List'),p.text
 results[kind]=dict(paragraphs=len(d.paragraphs),tables=len(d.tables),images=len(d.inline_shapes),equations=len(d._element.xpath('.//m:oMath')),words=len(text.split()),sha256=hashlib.sha256((docs/f'CG_{kind}_ReplayVerified.docx').read_bytes()).hexdigest())
(r/'final_content_verification.json').write_text(json.dumps(results,indent=2));print(json.dumps(results,indent=2))
