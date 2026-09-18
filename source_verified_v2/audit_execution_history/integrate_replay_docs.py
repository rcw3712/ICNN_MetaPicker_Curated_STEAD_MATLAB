from pathlib import Path
import csv,json,copy,re,sys,hashlib
from docx import Document
from lxml import etree
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

w=Path(__file__).parent;r=w/'audit_replay_threshold_station_20260918';n=w/'audit_no_training_20260918'
src=Path(r'C:\Drive E\Publikasi\2026-zd Computer & Geosciences\Submission\CG_Revision_20260918_SourceVerified\Documents')
out=r/'Documents';out.mkdir(exist_ok=True)
v=json.loads((r/'replay_verification.json').read_text());assert v['passed']
tv=json.loads((r/'threshold/verification.json').read_text());assert tv['primary_paired_bootstrap_reproduced']
def read(p):
 with Path(p).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
boot=read(r/'threshold/threshold_paired_bootstrap.csv')
def replace(p,text):
 fmt=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();run=p.add_run(text)
 if fmt is not None:run._r.insert(0,fmt)
def after(p,text):
 from docx.text.paragraph import Paragraph
 el=OxmlElement('w:p');p._p.addnext(el);np=Paragraph(el,p._parent);np.style=p.style
 np.add_run(text);np.paragraph_format.line_spacing=2
 for run in np.runs:run.font.name='Times New Roman';run.font.size=Pt(12)
 return np
def range_rows(policy,phase,comp):return [x for x in boot if x['Policy']==policy and x['Phase']==phase and x['Comparison']==comp]
def rng(policy,phase,comp):
 rows=range_rows(policy,phase,comp);a=[float(x['Delta']) for x in rows]
 return f'{min(a):.4f}–{max(a):.4f}'
def ci_count(policy,phase,comp,above=True):
 return sum(float(x['CI_low'])>0 if above else float(x['CI_high'])<0 for x in range_rows(policy,phase,comp))
sc='Stack Full3C minus Zonly';pc='PhaseNet Full3C minus Zonly'
threshold_result=(f'Across the predefined decoder grid, accepted S Full3C-minus-Z-only ΔF1 ranged from {rng("accepted","S",sc)} for stacking and {rng("accepted","S",pc)} for PhaseNetMatched. '
 f'Pointwise source-cluster intervals were above zero in {ci_count("accepted","S",sc)}/27 and {ci_count("accepted","S",pc)}/27 configurations, respectively. '
 'Quality-threshold changes leave accepted scoring unchanged by construction; the 27 settings therefore reduce to nine peak/window combinations for accepted scoring, not 27 independent checks. '
 f'Detected-only S differences ranged from {rng("strict","S",sc)} and {rng("strict","S",pc)}, respectively. '
 'The PhaseNetMatched P advantage retained intervals below zero in both modes at every setting. Accepted S method-comparison intervals included zero throughout, but detected-only Z-only S favoured the baseline in 18/27 pointwise intervals. The stacking P component difference had intervals above zero in 9/27 accepted and 27/27 detected-only configurations, showing dependence on the reporting policy. Supplementary Table S27 reports every contrast summary and Fig. S7 shows the quality-3 slice. These exploratory results neither select a new operating point nor establish historical calibration provenance.')
replay_detail=(f'Fresh GPU inference replay loaded all 106 v2 checkpoints and rebuilt features from the original waveform CSVs. Across six mode/base combinations, all 18,528 original and augmented OOF tensors and 1,980 test meta tensors agreed with their saved counterparts within absolute tolerance 1e-6 '
 f'(maximum differences {v["max_oof_feature_difference"]:.3g} and {v["max_test_feature_difference"]:.3g}, respectively). '
 f'All 70 final prediction sets, each with 330 records, reproduced constrained and unconstrained pick availability, statuses and times within 1e-7 s, with zero mismatches. '
 'Saved validation-derived ensemble weights were retained. This verifies inference and feature reconstruction for the frozen artifacts, not retraining determinism or independent predictive validation.')

main=Document(src/'CG_MainText_SourceVerified.docx');p=main.paragraphs
updates={
11:'Source overlap can compromise evaluation of seismic phase pickers and stacked models. We evaluate temporal stacking of STA/LTA, AIC, CNN and TCN outputs with waveform context using recovered STEAD earthquake identities. An exhaustive local-archive fingerprint audit uniquely matches 2,234 retained records from 1,477 sources. A fixed source-disjoint split assigns 330 records from 221 sources to testing, although historical development exposure prevents an independent-confirmation claim. Source-grouped out-of-fold predictions supply meta-training features. Three base seeds crossed with three meta seeds yielded Full3C accepted-pick mean F1 at ±100 ms of 0.8268 for P and 0.5536 for S, versus 0.8134 and 0.2686 for Z-only. The S difference was 0.2850, with a conditional paired source-cluster 95% interval of 0.2293–0.3418. Restricting evaluation to 287 records with both arrivals marked manual retained an S difference of 0.2705 (0.2113–0.3299). A PhaseNet-style MATLAB baseline achieved Full3C P/S F1 of 0.9833/0.5857 and Z-only F1 of 0.9833/0.2995. Conditional intervals supported its P advantage in both modes, whereas both S method-comparison intervals included zero. Fresh inference replay verifies all 106 checkpoints, OOF features and 70 test prediction sets. The contribution is an inspectable workflow linking waveform identity, grouped feature generation, checkpoint provenance and evaluation. Historical data reuse, one fixed split and limited unseen-station observations constrain generalization.',
19:p[19].text.replace('The standalone baseline has higher mean F1','Under the primary acceptance policy, the standalone baseline has higher mean F1'),
20:p[20].text.replace('unseen earthquake sources within the curated population','earthquake sources held out from corrected fitting and validation within the retained population'),
34:'Each 6,000×3 CSV waveform was converted to little-endian float32 in column-major E/N/Z order and matched by SHA-256 fingerprint to upstream STEAD HDF5 arrays. An exhaustive audit subsequently hashed all 1,265,657 waveform datasets in the local archive, including noise records without metadata or arrival filtering. All 2,234 selected fingerprints were distinct and each matched exactly one trace; no ambiguous trace or conflicting source assignment was found. Every recovered trace/source mapping agreed with upstream metadata. The mapping, numerical fingerprints, source partitions and OOF memberships were frozen before v2 training; the exhaustive uniqueness audit followed training. This establishes exact numerical uniqueness in this local archive, not near-duplicate absence or correctness of every upstream source identifier.',
39:p[39].text+' Historical roles remain relevant after correction: 233 of the current 330 test records formerly belonged to training, 46 to validation and 51 to testing. At source level, 196 of 221 current test sources had at least one record in historical training or validation. The current source-disjoint run is therefore a retrospective evaluation, not a development-naive confirmation (Supplementary Table S23).',
49:p[49].text+' Both-phase manual status applies to 1,224/1,544 training, 306/360 validation and 287/330 test records. A manual-reference sensitivity analysis restricts evaluation only; it does not retrain on manual labels or establish independent label accuracy (Supplementary Tables S24–S25).',
57:p[57].text+' Of the 64 test network–receiver keys, 55 also occur in training and account for 318/330 test records. The nine keys absent from both training and validation contain only 12 test records from ten sources. Supplementary Table S26 reports this small retrospective stratum; the experiment was not designed as a station holdout.',
62:p[62].text+' Supplementary Table S9 enumerates the 40 main invocations; Table S18 accounts for the full 106 fits.',
66:p[66].text.replace('Test sources were excluded from fitting and checkpoint selection.','Current test sources were excluded from v2 fitting and checkpoint selection; this exclusion does not erase their historical roles before identity correction.'),
82:p[82].text+' A 27-setting sensitivity grid was specified before its calculation: peak threshold {0.20, 0.30, 0.40}, tied P/S quality threshold {2, 3, 4}, and maximum S–P window {20, 30, 40} s, with minimum 0.1 s and uncertain boundary half the peak threshold. All settings are reported for the 18 full stackers and six baseline fits; the original operating point remains primary. Window changes alter both the S search domain and its quality-score denominator.',
90:p[90].text+' Subset diagnostics resample the sources within each stated subset. These intervals do not separately quantify station-level sampling uncertainty. Threshold-grid intervals are pointwise and exploratory; their envelope is not a simultaneous confidence interval, and no test-grid configuration is selected.',
94:p[94].text.replace('An independent numerical audit checks','Fresh inference replay and an independent numerical audit check'),
95:replay_detail,
98:p[98].text.replace('Supplementary Tables S15–S18','Supplementary Tables S9 and S15–S18'),
112:'No recovered earthquake source is shared across the current train, validation and test subsets. This excludes current fitting and checkpoint-selection overlap, but not historical development exposure. Only 25 current test sources (26 records) lack historical training or validation membership; all were previously in the historical test set. No development-naive subset has been established. The earlier partition used incorrect identities and is not a valid source-held-out reference. Because the corrected run also changes population allocation and training, no controlled estimate of leakage-induced score inflation is made.',
140:'The real-data verification now links 2,234 unique local-archive waveform matches to source-disjoint model memberships, fresh inference and numerical summaries. All 106 checkpoints pass source-exclusion checks; six actual OOF manifests match the frozen protocol. Fresh reconstruction covers 18,528 OOF tensors, 1,980 test meta tensors and all 70 prediction sets under both decoders. Independently recomputed metrics reproduce all 1,260 rows, and 25,248 manifest entries agree with 2,336 unique input files (Supplementary Table S21). These are computational reproducibility checks on fixed artifacts. They do not convert historically reused observations into an independent test population.',
147:p[147].text+' The historical role audit shows that most current test sources previously occurred in training or validation. Manual-reference sensitivity does not remove development exposure or quantify annotation error. Only 12 records occur at network–receiver keys absent from training and validation, precluding a strong station-transfer conclusion. Decoder sensitivity does not retrospectively establish validation-only threshold calibration.',
154:'Further validation requires observations held outside the preceding research and model-development process. A new random split of this same population would not by itself remove historical exposure. Additional source-grouped splits can quantify partition sensitivity, while nested validation can separate model-selection stages. Independent stations and datasets, official picker implementations, controlled input and compute budgets, continuous-record false alarms and operational latency remain appropriate next tests. Controlled perturbations and comparators isolating temporal context from nonlinearity would clarify robustness and mechanism.',
156:'The corrected workflow connects uniquely matched local STEAD waveform identities, source-disjoint outer and OOF partitions, checkpoint memberships, fresh inference replay and independently reproducible metrics. Across three base seeds crossed with three meta seeds, the Full3C–Z-only S difference persisted on one fixed split and in the both-phase manual-reference sensitivity analysis. This supports a component-access finding within the retained population. Historical development exposure remains, and station, regional and independent-dataset generalization are not established.',
157:p[157].text.replace('221 unseen sources','221 sources held out from v2 fitting and validation'),
158:p[158].text.replace('while both S method-comparison intervals included zero','while both primary S method-comparison intervals included zero').replace('Further source splits and independent datasets are required before broader robustness claims.','Independent data kept outside prior development are needed for confirmation; additional splits of the retained population would quantify partition sensitivity only.'),
162:p[162].text.replace('independent audit outputs','independent audit outputs, fresh inference receipts and threshold/station diagnostics'),
165:p[165].text+' The exhaustive fingerprint index verifies exact-match multiplicity in the local HDF5 version; it does not establish uniqueness across other STEAD versions. Historical role transitions and all threshold/station sensitivity outputs are retained with the audit package.'}
for i,t in updates.items():replace(p[i],t)
np=after(p[117],'Manual-reference evaluation retained 287 records from 188 sources. The expanded stacking S Full3C-minus-Z-only difference was 0.2705 (conditional 95% interval 0.2113–0.3299), and all nine paired fits remained positive. PhaseNetMatched gave 0.2550 (0.2013–0.3085). This supports consistency within the manual-status subset, not label-error immunity. In the 20-record intersection of manual references and no historical training/validation membership, both methods’ intervals included zero; that subset is small and was historically tested (Supplementary Table S25; Fig. S6).')
np=after(np,'Station stratification is also descriptive. The 12 records at nine network–receiver keys unseen in current training and validation yielded accepted S F1 of 0.8333/0.2593 for Full3C/Z-only stacking and 0.8889/0.2222 for PhaseNetMatched. Only ten sources contribute; the method-comparison intervals include zero. These observations cannot establish station generalization (Supplementary Table S26).')
after(np,threshold_result)
main.save(out/'CG_MainText_ReplayVerified.docx')

def style_table(t,widths=None):
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for col,width in zip(t.columns,widths):col.width=Inches(width)
 pr=t._tbl.tblPr
 borders=pr.find(qn('w:tblBorders'))
 if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
 for child in list(borders):borders.remove(child)
 for tag in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 for ir,row in enumerate(t.rows):
  if row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is None:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if ir==0 and row._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr()
   for oldpr in list(tcpr):
    if oldpr.tag in [qn('w:shd'),qn('w:tcMar')]:tcpr.remove(oldpr)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6' if ir==0 else 'FFFFFF');tcpr.append(sh)
   margin=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'70');x.set(qn('w:type'),'dxa');margin.append(x)
   tcpr.append(margin)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=1;para.paragraph_format.keep_with_next=(ir==0)
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(9);run.bold=ir==0
def fill(t,rows):
 while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
 for i,row in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for c,value in zip(cells,row):c.text=str(value)
 style_table(t)
def addtable(d,number,caption,rows,widths):
 cp=d.add_paragraph(f'Supplementary Table S{number}. {caption}');cp.paragraph_format.keep_with_next=True
 t=d.add_table(rows=1,cols=len(rows[0]));fill(t,rows);style_table(t,widths)
 return t
supp=Document(src/'CG_Supplementary_SourceVerified.docx');p=supp.paragraphs
for i,text in {
3:'Scope of results. Empirical tables and figures use the corrected source-verified v2 run. Base-42 diagnostics vary meta seeds 42–44; expanded summaries cross base seeds 42–44 with those meta seeds. Fresh inference replay verifies the current artifacts. Historical data roles, reference-label sensitivity and station/threshold diagnostics qualify predictive interpretation.',
24:p[24].text.replace('Test labels are used for scoring, not weight fitting or checkpoint selection.','Current test labels are used for scoring, not v2 weight fitting or checkpoint selection; historical roles before correction are documented in Table S23.'),
33:p[33].text+' Fifty-five of 64 test network–receiver keys occur in training (318 records). Nine unseen keys contain 12 records from ten sources and are also absent from validation. Table S26 gives retrospective performance strata; these keys do not identify station epochs.',
52:p[52].text+' Table S27 reports a predefined sensitivity grid applied to fresh fixed-model predictions, without selecting a replacement threshold.',
65:p[65].text+' Fresh waveform-to-inference replay is detailed in Table S21.',
70:p[70].text+' The subsequent 27-setting sensitivity analysis retains the primary configuration and does not resolve the unknown initial calibration history.',
71:'Supplementary Table S21. Current numerical, source-verification and fresh inference replay scope and outcomes.',
72:replay_detail+' Replay reads the original CSV time grids and waveforms, reapplies conditioning and each input mode, and reconstructs original and augmented OOF inputs in saved order with augmentation seed 4242. Each OOF model excludes its held source from fitting and inner validation. Checkpoint hashes are checked before loading, and all code/model hashes are checked again after completion. Frozen ensemble weights are not re-estimated. Prediction receipts separately report pick/status mismatch counts and numerical quality-score differences. No training is performed.',
75:'The exhaustive audit hashed all 1,265,657 waveform datasets in the local HDF5 archive, with no phase-label or metadata filtering. All 2,234 selected fingerprints were distinct and each matched exactly one trace, with zero conflicting source assignments or read errors. Hashes use little-endian float32 values in column-major E/N/Z order, without normalization. Current CSV fingerprints independently agree with the frozen mapping. This verifies exact numerical multiplicity in this archive version, not near duplicates or catalogue identity truth. Initial identity recovery and split freezing preceded v2 training; exhaustive uniqueness verification followed training. The original export recipe remains unavailable. All CSV P/S labels equal floor(upstream arrival sample)/100; 94 P and 124 S labels differ from fractional upstream samples by at most 9.01 ms. These labels were retained consistently for all corrected training and evaluation.'}.items():replace(p[i],text)
for row in supp.tables[18].rows:
 if row.cells[0].text=='Model replay':row.cells[2].text='Fresh v2 replay: all 106 checkpoints; test predictions and OOF tensors (Table S21)'
rows=[[c.text for c in row.cells] for row in supp.tables[20].rows][:-1]
rows += [
 ['Fresh test inference','70 sets × 330 records; two decoders','Zero pick/status mismatches; times within 1e-7 s'],
 ['Fresh test meta features','6 mode/base groups × 330 tensors',f'Maximum absolute difference {v["max_test_feature_difference"]:.3g}; tolerance 1e-6'],
 ['Fresh OOF features','6 groups × 3088 original/augmented tensors',f'Maximum absolute difference {v["max_oof_feature_difference"]:.3g}; tolerance 1e-6'],
 ['Replay integrity','106 checkpoint hashes and 89 code files','Unchanged after replay; no training']]
fill(supp.tables[20],rows)
rows=[[c.text for c in row.cells] for row in supp.tables[21].rows]
rows += [['Exhaustive fingerprint multiplicity','1,265,657 local waveform datasets','2234 distinct selected hashes; one match each'],['Ambiguous identity / read failures','All selected matches / full archive','0 conflicting sources / 0 read errors']]
fill(supp.tables[21],rows)

supp.add_paragraph('Retrospective sensitivity and provenance',style='Heading 1')
addtable(supp,23,'Historical-to-current role transitions. Record counts are exclusive; source counts indicate any membership and are nonexclusive.',[
 ['Unit','Historical role','Current train','Current validation','Current test'],
 ['Record','Train',1072,251,233],['Record','Validation',249,48,46],['Record','Test',223,61,51],
 ['Source','Train',804,177,177],['Source','Validation',225,46,44],['Source','Test',209,53,48]],[.75,1.1,1.1,1.25,1.1])
supp.add_paragraph('A physical source could occupy several historical roles, so source rows must not be summed. Of 221 current test sources, 196 had at least one historical training/validation record, covering 304 current test records. The remaining 25 sources (26 records) had historical test exposure. This role audit does not reconstruct every research decision or prove that every exposed record influenced one. It establishes no development-naive test subset. Exclusive source-role patterns and record-level transitions are supplied as CSVs.')
addtable(supp,24,'Recovered upstream reference-pick status by current split.',[
 ['Split','All records','Both phases manual','At least one nonmanual'],['Train',1544,1224,320],['Validation',360,306,54],['Test',330,287,43]],[1.2,1.1,1.55,1.55])
supp.add_paragraph('Nonmanual denotes automatic or autopicker status for at least one arrival. Manual status is a metadata category, not an independent reannotation. The manual-only test sensitivity retains 287 records from 188 sources; the complement contains 43 records from 36 sources. Some sources contribute to both strata. All original model fits are retained.')
addtable(supp,25,'S component-access sensitivity at ±100 ms under accepted scoring. Values average nine stackers or three baseline fits; differences are calculated before rounding. Conditional paired 95% intervals use 10,000 source-cluster draws.',[
 ['Evaluation subset','Method','N / sources','Full3C F1','Z-only F1','ΔF1 [95% interval]'],
 ['All test','Stack','330 / 221','.5536','.2686','.2850 [.2293, .3418]'],
 ['Both manual','Stack','287 / 188','.5476','.2770','.2705 [.2113, .3299]'],
 ['Both manual','PhaseNet','287 / 188','.5655','.3105','.2550 [.2013, .3085]'],
 ['No historical train/val','Stack','26 / 25','.4444','.1974','.2470 [.0582, .4462]'],
 ['No historical train/val','PhaseNet','26 / 25','.4643','.2451','.2192 [.0503, .3905]'],
 ['Manual and no historical train/val','Stack','20 / 19','.3647','.1599','.2047 [−.0092, .4390]'],
 ['Manual and no historical train/val','PhaseNet','20 / 19','.3333','.2342','.0991 [−.0794, .2802]']],[1.35,.7,.7,.7,.7,1.35])
supp.add_paragraph('Manual-subset stacking P F1 is 0.8093/0.7982 for Full3C/Z-only; PhaseNetMatched gives 0.9913/0.9843. Its P advantage remains clear in conditional intervals. Manual S stacker-minus-baseline intervals include zero in both modes. All nine stacking component differences remain positive for manual S (range 0.2403–0.3078). The smallest intersection yields intervals spanning zero for both methods. Neither the subset comparisons nor their overlapping populations isolate a causal annotation effect; the absence of historical train/validation membership does not remove historical test exposure. Full per-fit metrics and intervals for six subsets accompany this table.')

st=read(r/'station/sensitivity_summary.csv')
# Match schema dynamically but require exact subset names from the exported audit.
station_meta=json.loads((r/'station/station_summary.json').read_text())
stationrows=[['Station-key stratum','Method','N / sources','Full3C P / S','Z-only P / S']]
def stval(subset,method,mode,phase):
 rows=[x for x in st if x['subset']==subset and x['Method']==method and x['Mode']==mode and x['Phase']==phase and x['Policy']=='accepted' and x['Tolerance_ms']=='100']
 assert len(rows)==1,(subset,method,mode,phase,st[0]);return float(rows[0]['mean_F1'])
# Actual exported subset identifiers are inspected and supplied explicitly.
for subset,label,nn in [('station_seen_train','Seen in training','318 / 211'),('station_unseen_train','Unseen in training','12 / 10')]:
 for method,labelmethod in [('Stack','Stack'),('PhaseNetMatched','PhaseNet')]:
  stationrows.append([label,labelmethod,nn]+[f'{stval(subset,method,mode,"P"):.4f} / {stval(subset,method,mode,"S"):.4f}' for mode in ['Full3C','Zonly']])
addtable(supp,26,'Station-key sensitivity at ±100 ms with accepted scoring. Network and receiver codes identify keys, not station epochs.',stationrows,[1.35,.8,1.0,1.15,1.15])
supp.add_paragraph('The seen stratum contains 55 test keys. The nine keys unseen in training are also absent from validation and include only 12 records from ten sources. This post hoc stratum is not a station-held-out sampling design. For unseen-key S, Full3C-minus-Z-only intervals are [0.3131, 0.7863] for stacking and [0.3636, 0.8974] for PhaseNetMatched. Stacker-minus-baseline S intervals are [−0.1667, 0.0667] in Full3C and [−0.0684, 0.1444] in Z-only. These intervals resample ten sources and do not separately quantify station-level sampling uncertainty. Such a small conditional diagnostic does not support a general station-transfer claim.')

thresholdrows=[['Policy / phase','Contrast','Minimum ΔF1','Maximum ΔF1','CI above / below zero']]
names={sc:'Stack Full3C − Z',pc:'PhaseNet Full3C − Z','Stack minus PhaseNet Full3C':'Stack − PhaseNet Full3C','Stack minus PhaseNet Zonly':'Stack − PhaseNet Z'}
for policy in ['accepted','strict']:
 for phase in ['P','S']:
  for comp,label in names.items():
   rr=range_rows(policy,phase,comp);vals=[float(x['Delta']) for x in rr]
   thresholdrows.append([f'{policy} / {phase}',label,f'{min(vals):.4f}',f'{max(vals):.4f}',f'{ci_count(policy,phase,comp)} / {ci_count(policy,phase,comp,False)} of 27'])
addtable(supp,27,'Predefined threshold-grid sensitivity at ±100 ms. Ranges span 27 settings; counts refer to pointwise exploratory intervals.',thresholdrows,[.95,1.65,.85,.85,1.15])
supp.add_paragraph('The Cartesian grid varies peak threshold 0.20/0.30/0.40, tied P/S quality threshold 2/3/4 and maximum S–P window 20/30/40 s. Minimum separation remains 0.1 s and the uncertain boundary is half the peak threshold. Eighteen full stacker and six baseline fits are decoded at every setting; the published 0.30/3/30 operating point is unchanged. All configurations, both scoring policies, per-fit metrics and conditional source-bootstrap intervals are supplied as CSVs. Counts are not multiplicity-adjusted tests. Accepted scoring combines detected and uncertain statuses, so quality changes cannot alter its mask or S gate; it reduces to nine peak/window combinations. Strict scoring uses detected candidates but still gates S using accepted P. Changing the S window also changes its quality-score mean. Figure S7 displays only the quality-3 slice; the table and files cover all quality values.')
supp.add_paragraph(threshold_result)
for num,path,caption in [
 (6,r/'Figures/S6_roles_manual_sensitivity.png','S component-access sensitivity by reference status and historical-role restriction. Source-cluster intervals condition on fixed fits. No subset shown is established to be development-naive.'),
 (7,r/'Figures/S7_threshold_sensitivity.png','S Full3C-minus-Z-only F1 sensitivity at ±100 ms for the quality-threshold-3 slice. Columns vary peak threshold; rows vary maximum S–P window. The original configuration is outlined. Table S27 and accompanying CSVs include all three quality thresholds. No operating point is selected from these results.')]:
 cp=supp.add_paragraph(f'Supplementary Fig. S{num}. {caption}');cp.paragraph_format.keep_with_next=True
 pp=supp.add_paragraph();pp.add_run().add_picture(str(path),width=Inches(6.0));pp.paragraph_format.keep_with_next=False
supp.save(out/'CG_Supplementary_ReplayVerified.docx')

cover=Document(src/'CG_Cover_Letter_SourceVerified.docx');p=cover.paragraphs
replace(p[7],p[7].text.replace('S method differences remain unresolved','primary S method differences remain unresolved').replace('stacking P/S F1','stacking accepted-pick P/S F1').replace('221 unseen sources','221 sources held out from v2 fitting and validation').replace('The persistent S component-access difference is shared by both approaches.','Both approaches show an S component-access difference on this split.'))
replace(p[8],'The computational contribution connects unique waveform identity, grouped OOF construction, feature interfaces and model provenance. Fresh replay verifies all 106 checkpoints, OOF features and 70 test prediction sets; independent calculations reproduce all 1,260 metric rows. Manual-reference, threshold and station diagnostics qualify the findings. Established grouping and stacking methods are not presented as new algorithms. Historical data reuse prevents an independent-confirmation claim; the small unseen-station subset does not establish transfer. The original extraction recipe remains unavailable.')
cover.save(out/'CG_Cover_Letter_ReplayVerified.docx')
high=Document(src/'CG_Highlights_SourceVerified.docx')
highlights=['An exhaustive STEAD fingerprint audit uniquely matches all 2234 retained records.','Replay verifies OOF features and test predictions from 106 saved checkpoints.','Manual-reference sensitivity retains the Full3C advantage for S picking.','PhaseNet-style has higher P F1; primary S method intervals include zero.','Historical data reuse limits claims of independent validation.']
for p,t in zip(high.paragraphs[1:],highlights):assert len(t)<=85;replace(p,t)
high.save(out/'CG_Highlights_ReplayVerified.docx')

checks={'abstract_words':len(updates[11].split()),'highlight_characters':[len(x) for x in highlights],'files':[]}
assert checks['abstract_words']<=250
for file in out.glob('*.docx'):
 d=Document(file);old=Document(src/file.name.replace('ReplayVerified','SourceVerified'))
 assert len(d._element.xpath('.//m:oMath'))==len(old._element.xpath('.//m:oMath'))
 if 'MainText' in file.name:
  assert [etree.tostring(x, method="c14n") for x in d._element.xpath('.//m:oMath')]==[etree.tostring(x, method="c14n") for x in old._element.xpath('.//m:oMath')]
 checks['files'].append(dict(name=file.name,tables=len(d.tables),images=len(d.inline_shapes),equations=len(d._element.xpath('.//m:oMath')),sha256=hashlib.sha256(file.read_bytes()).hexdigest()))
(r/'document_content_checks.json').write_text(json.dumps(checks,indent=2));print(json.dumps(checks,indent=2))
