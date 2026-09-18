from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
w=Path(__file__).parent;r=w/'audit_replay_threshold_station_20260918';p=r/'Documents/CG_MainText_ReplayVerified.docx';d=Document(p)
for i in [39,92,94,95,134,136,143,144]:
 para=d.paragraphs[i];para.paragraph_format.line_spacing=2
 for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(12)
for i,style in [(93,'Heading 2'),(142,'Heading 3')]:
 para=d.paragraphs[i];para.style=style;para.paragraph_format.keep_with_next=True
 for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(12);run.font.color.rgb=RGBColor(0,0,0)
d.save(p)
